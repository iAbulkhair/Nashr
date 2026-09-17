"""Build Nashr 36-month financial model (SAR) with live formulas."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = "docs/Nashr_Financial_Model.xlsx"
MONTHS = 36
START = (2026, 10)

wb = Workbook()
NAVY = "0F3D3A"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=NAVY)
inp_fill = PatternFill("solid", fgColor="FFF4D6")
sec_font = Font(bold=True, size=12, color=NAVY)
bold = Font(bold=True)
thin = Side(style="thin", color="D9D9D9")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
SAR = '#,##0;[Red](#,##0)'
PCT = '0.0%'
DEC = '0.00'


def month_label(i):
    y, m = START
    m += i
    y += (m - 1) // 12
    m = (m - 1) % 12 + 1
    return f"{['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'][m-1]}-{str(y)[2:]}"


def header(ws, row, labels, start_col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + j, value=lab)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box


A = wb.active
A.title = "Assumptions"
A.column_dimensions["A"].width = 44
A.column_dimensions["B"].width = 16
A.column_dimensions["C"].width = 60
A["A1"] = "Nashr — Model Assumptions (SAR)"
A["A1"].font = Font(bold=True, size=14, color=NAVY)
A["A2"] = "Yellow cells are inputs. Everything else is formula-driven."
A["A2"].font = Font(italic=True, color="666666")
names = {}
r = 4


def section(title):
    global r
    r += 1
    A.cell(row=r, column=1, value=title).font = sec_font
    r += 1


def inp(name, label, value, fmt=None, note=""):
    global r
    A.cell(row=r, column=1, value=label)
    c = A.cell(row=r, column=2, value=value)
    c.fill = inp_fill
    c.border = box
    if fmt:
        c.number_format = fmt
    A.cell(row=r, column=3, value=note).font = Font(color="666666", italic=True)
    names[name] = f"Assumptions!$B${r}"
    r += 1


section("Pricing (monthly, SAR, ex-VAT)")
inp("p_starter", "Starter price", 149, SAR)
inp("p_pro", "Pro price", 399, SAR)
inp("p_agency", "Agency price", 799, SAR)
inp("p_ent", "Enterprise avg contract / month", 4500, SAR, "Custom deals; blended average")
inp("annual_disc", "Annual plan discount (2 months free)", 2 / 12, PCT)
inp("annual_mix", "Share of customers on annual", 0.30, PCT)
inp("topup_pct", "Top-up revenue as % of subscription", 0.12, PCT, "Credit top-ups + extra seats/brands")
section("Plan mix of paying customers")
inp("mix_starter", "Starter", 0.45, PCT)
inp("mix_pro", "Pro", 0.40, PCT)
inp("mix_agency", "Agency", 0.13, PCT)
inp("mix_ent", "Enterprise", 0.02, PCT)
section("Acquisition")
inp("trials_m1", "Trials in month 1", 180, "#,##0", "Waitlist + design partners")
inp("trial_growth_y1", "Monthly trial growth, year 1", 0.24, PCT)
inp("trial_growth_y2", "Monthly trial growth, year 2", 0.09, PCT)
inp("trial_growth_y3", "Monthly trial growth, year 3", 0.03, PCT)
inp("conv", "Trial → paid conversion", 0.12, PCT)
inp("churn_m", "Monthly logo churn", 0.04, PCT, "Target <4%; SMB SaaS median 3–6%")
inp("cac", "Blended CAC per paying customer", 650, SAR, "Salla/Zid app store, content, paid, agency rev-share")
section("Cost of revenue (per paying brand per month)")
inp("llm_starter", "AI/model + media cost — Starter", 22, SAR, "200 credits × ~0.11 SAR")
inp("llm_pro", "AI/model + media cost — Pro", 70, SAR, "700 credits × ~0.10 SAR")
inp("llm_agency", "AI/model + media cost — Agency", 190, SAR, "2,000 credits × ~0.095 SAR")
inp("llm_ent", "AI/model + media cost — Enterprise", 900, SAR)
inp("infra_per_brand", "Hosting/infra per brand (KSA region)", 6, SAR)
inp("msg_per_brand", "WhatsApp/SMS pass-through per brand", 9, SAR, "Cloud API conversation fees, avg")
inp("pay_fee", "Payment processing % of revenue", 0.028, PCT, "Tap/Moyasar mada+cards blend")
inp("support_per_brand", "Support/QA cost per brand", 8, SAR, "Arabic QA reviewers, tier-1 support")
section("Operating expenses")
inp("tools_m", "Software & tools / month", 9000, SAR)
inp("office_m", "Office & admin / month", 12000, SAR)
inp("legal_m", "Legal, accounting, compliance / month", 7000, SAR)
inp("mkt_base", "Marketing base (content, events) / month", 15000, SAR, "Excludes CAC-driven spend (new paying × CAC)")
inp("vat", "VAT (pass-through on collections; not in P&L)", 0.15, PCT)
section("Headcount (SAR fully loaded per month)")
inp("sal_founder", "Founder / CEO", 25000, SAR)
inp("sal_eng", "Engineer", 30000, SAR)
inp("sal_ai", "AI / prompt engineer", 32000, SAR)
inp("sal_design", "Designer", 18000, SAR)
inp("sal_growth", "Growth / partnerships", 22000, SAR)
inp("sal_cs", "Customer success / Arabic QA", 12000, SAR)
inp("sal_sales", "Enterprise sales", 24000, SAR)
section("Funding")
inp("cash0", "Opening cash", 500000, SAR, "Founder + angels")
inp("raise_amt", "Seed raise", 6000000, SAR, "~USD 1.6M")
inp("raise_month", "Seed raise month (1–36)", 3, "0")

H = wb.create_sheet("Headcount")
H.column_dimensions["A"].width = 30
roles = [
    ("Founder / CEO", "sal_founder", [1] * 36),
    ("Engineer", "sal_eng", [1] * 3 + [2] * 5 + [3] * 6 + [4] * 10 + [6] * 6 + [8] * 6),
    ("AI / prompt engineer", "sal_ai", [1] * 6 + [2] * 30),
    ("Designer", "sal_design", [0] * 4 + [1] * 32),
    ("Growth / partnerships", "sal_growth", [0] * 3 + [1] * 12 + [2] * 12 + [3] * 9),
    ("Customer success / Arabic QA", "sal_cs", [1] * 4 + [2] * 6 + [3] * 8 + [5] * 6 + [8] * 6 + [12] * 6),
    ("Enterprise sales", "sal_sales", [0] * 9 + [1] * 12 + [2] * 9 + [3] * 6),
]
header(H, 1, ["Role", "Cost/mo"] + [month_label(i) for i in range(MONTHS)])
for k, (role, sal, counts) in enumerate(roles):
    rr = 2 + k
    H.cell(row=rr, column=1, value=role)
    H.cell(row=rr, column=2, value=f"={names[sal]}").number_format = SAR
    for i, n in enumerate(counts):
        c = H.cell(row=rr, column=3 + i, value=n)
        c.fill = inp_fill
        c.border = box
tot_row = 2 + len(roles)
H.cell(row=tot_row, column=1, value="Total headcount").font = bold
H.cell(row=tot_row + 1, column=1, value="Total payroll (SAR)").font = bold
for i in range(MONTHS):
    col = L(3 + i)
    H.cell(row=tot_row, column=3 + i, value=f"=SUM({col}2:{col}{tot_row-1})").font = bold
    parts = "+".join(f"{col}{2+k}*$B${2+k}" for k in range(len(roles)))
    c = H.cell(row=tot_row + 1, column=3 + i, value=f"={parts}")
    c.number_format = SAR
    c.font = bold
H.freeze_panes = "C2"

C = wb.create_sheet("Customers")
C.column_dimensions["A"].width = 36
header(C, 1, ["Metric"] + [month_label(i) for i in range(MONTHS)])
rows = ["Month #", "Trials", "New paying", "Churned", "Ending paying brands", "— Starter", "— Pro", "— Agency", "— Enterprise"]
for k, lab in enumerate(rows):
    C.cell(row=2 + k, column=1, value=lab).font = bold if lab == "Ending paying brands" else Font()
R_MN, R_TR, R_NEW, R_CH, R_END, R_ST, R_PR, R_AG, R_EN = range(2, 11)
for i in range(MONTHS):
    col = L(2 + i)
    prev = L(1 + i)
    C.cell(row=R_MN, column=2 + i, value=i + 1)
    if i == 0:
        C.cell(row=R_TR, column=2 + i, value=f"={names['trials_m1']}")
    else:
        g = f"IF({col}{R_MN}<=12,{names['trial_growth_y1']},IF({col}{R_MN}<=24,{names['trial_growth_y2']},{names['trial_growth_y3']}))"
        C.cell(row=R_TR, column=2 + i, value=f"=ROUND({prev}{R_TR}*(1+{g}),0)")
    C.cell(row=R_NEW, column=2 + i, value=f"=ROUND({col}{R_TR}*{names['conv']},0)")
    if i == 0:
        C.cell(row=R_CH, column=2 + i, value=0)
        C.cell(row=R_END, column=2 + i, value=f"={col}{R_NEW}")
    else:
        C.cell(row=R_CH, column=2 + i, value=f"=ROUND({prev}{R_END}*{names['churn_m']},0)")
        C.cell(row=R_END, column=2 + i, value=f"={prev}{R_END}+{col}{R_NEW}-{col}{R_CH}")
    C.cell(row=R_ST, column=2 + i, value=f"=ROUND({col}{R_END}*{names['mix_starter']},0)")
    C.cell(row=R_PR, column=2 + i, value=f"=ROUND({col}{R_END}*{names['mix_pro']},0)")
    C.cell(row=R_AG, column=2 + i, value=f"=ROUND({col}{R_END}*{names['mix_agency']},0)")
    C.cell(row=R_EN, column=2 + i, value=f"={col}{R_END}-{col}{R_ST}-{col}{R_PR}-{col}{R_AG}")
    for rr in range(2, 11):
        C.cell(row=rr, column=2 + i).number_format = "#,##0"
    C.cell(row=R_END, column=2 + i).font = bold
C.freeze_panes = "B2"

P = wb.create_sheet("P&L")
P.column_dimensions["A"].width = 40
header(P, 1, ["SAR"] + [month_label(i) for i in range(MONTHS)] + ["Year 1", "Year 2", "Year 3"])
labels = [
    ("Subscription revenue — Starter", "st"), ("Subscription revenue — Pro", "pr"),
    ("Subscription revenue — Agency", "ag"), ("Subscription revenue — Enterprise", "en"),
    ("Top-ups, seats, brands", "top"), ("TOTAL REVENUE", "rev"), ("", None),
    ("AI model + media", "cogs_ai"), ("Infra (KSA region)", "cogs_inf"),
    ("WhatsApp/SMS pass-through", "cogs_msg"), ("Payment fees", "cogs_pay"),
    ("Support & Arabic QA", "cogs_sup"), ("TOTAL COGS", "cogs"),
    ("GROSS PROFIT", "gp"), ("Gross margin %", "gm"), ("", None),
    ("Payroll", "pay"), ("Acquisition spend (new paying × CAC)", "acq"),
    ("Marketing base", "mkt"), ("Software & tools", "tools"),
    ("Office & admin", "office"), ("Legal & compliance", "legal"),
    ("TOTAL OPEX", "opex"), ("EBITDA", "ebitda"), ("EBITDA margin %", "em"), ("", None),
    ("MRR (end of month)", "mrr"), ("ARR run-rate", "arr"), ("Cumulative EBITDA", "cum"),
]
PR = {}
for k, (lab, key) in enumerate(labels):
    rr = 2 + k
    c = P.cell(row=rr, column=1, value=lab)
    if lab.isupper() or key in ("gp", "ebitda", "mrr", "arr"):
        c.font = bold
    if key:
        PR[key] = rr
ann_eff = f"(1-{names['annual_mix']}*{names['annual_disc']})"
for i in range(MONTHS):
    col = L(2 + i)
    hcol = L(3 + i)

    def put(key, formula, fmt=SAR, b=False):
        c = P.cell(row=PR[key], column=2 + i, value=formula)
        c.number_format = fmt
        if b:
            c.font = bold

    put("st", f"=Customers!{col}{R_ST}*{names['p_starter']}*{ann_eff}")
    put("pr", f"=Customers!{col}{R_PR}*{names['p_pro']}*{ann_eff}")
    put("ag", f"=Customers!{col}{R_AG}*{names['p_agency']}*{ann_eff}")
    put("en", f"=Customers!{col}{R_EN}*{names['p_ent']}")
    put("top", f"=SUM({col}{PR['st']}:{col}{PR['en']})*{names['topup_pct']}")
    put("rev", f"=SUM({col}{PR['st']}:{col}{PR['top']})", b=True)
    put("cogs_ai", f"=Customers!{col}{R_ST}*{names['llm_starter']}+Customers!{col}{R_PR}*{names['llm_pro']}+Customers!{col}{R_AG}*{names['llm_agency']}+Customers!{col}{R_EN}*{names['llm_ent']}")
    put("cogs_inf", f"=Customers!{col}{R_END}*{names['infra_per_brand']}")
    put("cogs_msg", f"=Customers!{col}{R_END}*{names['msg_per_brand']}")
    put("cogs_pay", f"={col}{PR['rev']}*{names['pay_fee']}")
    put("cogs_sup", f"=Customers!{col}{R_END}*{names['support_per_brand']}")
    put("cogs", f"=SUM({col}{PR['cogs_ai']}:{col}{PR['cogs_sup']})", b=True)
    put("gp", f"={col}{PR['rev']}-{col}{PR['cogs']}", b=True)
    put("gm", f"=IF({col}{PR['rev']}=0,0,{col}{PR['gp']}/{col}{PR['rev']})", PCT)
    put("pay", f"=Headcount!{hcol}{tot_row+1}")
    put("acq", f"=Customers!{col}{R_NEW}*{names['cac']}")
    put("mkt", f"={names['mkt_base']}")
    put("tools", f"={names['tools_m']}")
    put("office", f"={names['office_m']}")
    put("legal", f"={names['legal_m']}")
    put("opex", f"=SUM({col}{PR['pay']}:{col}{PR['legal']})", b=True)
    put("ebitda", f"={col}{PR['gp']}-{col}{PR['opex']}", b=True)
    put("em", f"=IF({col}{PR['rev']}=0,0,{col}{PR['ebitda']}/{col}{PR['rev']})", PCT)
    put("mrr", f"={col}{PR['rev']}", b=True)
    put("arr", f"={col}{PR['mrr']}*12", b=True)
    put("cum", f"={col}{PR['ebitda']}" if i == 0 else f"={L(1+i)}{PR['cum']}+{col}{PR['ebitda']}")
for y in range(3):
    ycol = 2 + MONTHS + y
    c1, c2 = L(2 + y * 12), L(2 + y * 12 + 11)
    for key, rr in PR.items():
        if key in ("gm", "em"):
            top = "gp" if key == "gm" else "ebitda"
            f = f"=IF(SUM({c1}{PR['rev']}:{c2}{PR['rev']})=0,0,SUM({c1}{PR[top]}:{c2}{PR[top]})/SUM({c1}{PR['rev']}:{c2}{PR['rev']}))"
            fmt = PCT
        elif key in ("mrr", "arr", "cum"):
            f, fmt = f"={c2}{rr}", SAR
        else:
            f, fmt = f"=SUM({c1}{rr}:{c2}{rr})", SAR
        c = P.cell(row=rr, column=ycol, value=f)
        c.number_format = fmt
        c.font = bold
        c.fill = PatternFill("solid", fgColor="E8F0EE")
P.freeze_panes = "B2"

K = wb.create_sheet("Cash")
K.column_dimensions["A"].width = 30
header(K, 1, ["SAR"] + [month_label(i) for i in range(MONTHS)])
for rr, lab in enumerate(["Opening cash", "EBITDA", "Equity raised", "Closing cash", "Runway (months at current burn)"], start=2):
    K.cell(row=rr, column=1, value=lab).font = bold if lab == "Closing cash" else Font()
for i in range(MONTHS):
    col, prev = L(2 + i), L(1 + i)
    K.cell(row=2, column=2 + i, value=f"={names['cash0']}" if i == 0 else f"={prev}5").number_format = SAR
    K.cell(row=3, column=2 + i, value=f"='P&L'!{col}{PR['ebitda']}").number_format = SAR
    K.cell(row=4, column=2 + i, value=f"=IF({i+1}={names['raise_month']},{names['raise_amt']},0)").number_format = SAR
    c = K.cell(row=5, column=2 + i, value=f"={col}2+{col}3+{col}4")
    c.number_format = SAR
    c.font = bold
    K.cell(row=6, column=2 + i, value=f'=IF({col}3>=0,"∞",ROUND({col}5/-{col}3,1))')
K["A8"] = "Minimum cash balance (36m)"
K["B8"] = f"=MIN(B5:{L(1+MONTHS)}5)"
K["B8"].number_format = SAR
K["B8"].font = bold
K.freeze_panes = "B2"

U = wb.create_sheet("Unit Economics")
U.column_dimensions["A"].width = 40
for col in "BCDE":
    U.column_dimensions[col].width = 16
U["A1"] = "Per-plan unit economics (monthly, SAR)"
U["A1"].font = Font(bold=True, size=13, color=NAVY)
header(U, 3, ["", "Starter", "Pro", "Agency", "Enterprise"])
ue_rows = [
    ("Price (ex-VAT, blended for annual mix)", [f"={names['p_starter']}*{ann_eff}", f"={names['p_pro']}*{ann_eff}", f"={names['p_agency']}*{ann_eff}", f"={names['p_ent']}"]),
    ("+ Top-ups / seats", [f"={c}4*{names['topup_pct']}" for c in "BCDE"]),
    ("Revenue per brand", [f"={c}4+{c}5" for c in "BCDE"]),
    ("AI model + media", [f"={names['llm_starter']}", f"={names['llm_pro']}", f"={names['llm_agency']}", f"={names['llm_ent']}"]),
    ("Infra + messaging + support", [f"={names['infra_per_brand']}+{names['msg_per_brand']}+{names['support_per_brand']}"] * 4),
    ("Payment fees", [f"={c}6*{names['pay_fee']}" for c in "BCDE"]),
    ("Gross profit per brand", [f"={c}6-{c}7-{c}8-{c}9" for c in "BCDE"]),
    ("Gross margin %", [f"={c}10/{c}6" for c in "BCDE"]),
    ("Expected lifetime (months) = 1/churn", [f"=1/{names['churn_m']}"] * 4),
    ("LTV (gross profit × lifetime)", [f"={c}10*{c}12" for c in "BCDE"]),
    ("CAC", [f"={names['cac']}"] * 4),
    ("LTV : CAC", [f"={c}13/{c}14" for c in "BCDE"]),
    ("CAC payback (months)", [f"={c}14/{c}10" for c in "BCDE"]),
]
for k, (lab, fs) in enumerate(ue_rows):
    rr = 4 + k
    U.cell(row=rr, column=1, value=lab).font = bold if lab.startswith(("Revenue", "Gross profit", "LTV", "CAC payback")) else Font()
    for j, f in enumerate(fs):
        c = U.cell(row=rr, column=2 + j, value=f)
        c.number_format = PCT if "%" in lab else (DEC if ("LTV : CAC" in lab or "payback" in lab or "lifetime" in lab) else SAR)
U["A20"] = "Blended (weighted by plan mix)"
U["A20"].font = sec_font
mixes = [names["mix_starter"], names["mix_pro"], names["mix_agency"], names["mix_ent"]]


def blend(row):
    return "=" + "+".join(f"{L(2+j)}{row}*{mixes[j]}" for j in range(4))


for k, (lab, f, fmt) in enumerate([
    ("Blended revenue per brand", blend(6), SAR),
    ("Blended gross profit per brand", blend(10), SAR),
    ("Blended gross margin", "=B22/B21", PCT),
    ("Blended LTV", blend(13), SAR),
    ("Blended LTV:CAC", f"=B24/{names['cac']}", DEC),
    ("Blended CAC payback (months)", f"={names['cac']}/B22", DEC),
]):
    rr = 21 + k
    U.cell(row=rr, column=1, value=lab)
    c = U.cell(row=rr, column=2, value=f)
    c.number_format = fmt
    c.font = bold

S = wb.create_sheet("Summary", 0)
S.column_dimensions["A"].width = 38
for col in "BCD":
    S.column_dimensions[col].width = 18
S["A1"] = "Nashr — Financial Summary (SAR)"
S["A1"].font = Font(bold=True, size=14, color=NAVY)
S["A2"] = "36-month model. Change inputs on the Assumptions sheet; all sheets recalculate."
S["A2"].font = Font(italic=True, color="666666")
header(S, 4, ["", "Year 1", "Year 2", "Year 3"])
Y = [L(2 + MONTHS), L(3 + MONTHS), L(4 + MONTHS)]
lastcol = [L(13), L(25), L(37)]
summary = [
    ("Revenue", "rev", SAR), ("Gross profit", "gp", SAR), ("Gross margin", "gm", PCT),
    ("Total opex", "opex", SAR), ("EBITDA", "ebitda", SAR), ("EBITDA margin", "em", PCT),
    ("Ending MRR", "mrr", SAR), ("Ending ARR run-rate", "arr", SAR),
]
for k, (lab, key, fmt) in enumerate(summary):
    rr = 5 + k
    S.cell(row=rr, column=1, value=lab).font = bold if key in ("rev", "ebitda", "arr") else Font()
    for j in range(3):
        S.cell(row=rr, column=2 + j, value=f"='P&L'!{Y[j]}{PR[key]}").number_format = fmt
rr = 5 + len(summary)
S.cell(row=rr, column=1, value="Ending paying brands")
for j in range(3):
    S.cell(row=rr, column=2 + j, value=f"=Customers!{lastcol[j]}{R_END}").number_format = "#,##0"
rr += 1
S.cell(row=rr, column=1, value="Ending headcount")
for j in range(3):
    S.cell(row=rr, column=2 + j, value=f"=Headcount!{L(14 + j*12)}{tot_row}").number_format = "#,##0"
rr += 1
S.cell(row=rr, column=1, value="Ending cash")
for j in range(3):
    S.cell(row=rr, column=2 + j, value=f"=Cash!{lastcol[j]}5").number_format = SAR
rr += 2
for lab, f, fmt in [
    ("Blended LTV:CAC", "='Unit Economics'!B25", DEC),
    ("Blended CAC payback (months)", "='Unit Economics'!B26", DEC),
    ("Minimum cash balance (36m)", "=Cash!B8", SAR),
    ("Seed raise", f"={names['raise_amt']}", SAR),
]:
    S.cell(row=rr, column=1, value=lab).font = bold
    S.cell(row=rr, column=2, value=f).number_format = fmt
    rr += 1

wb.save(OUT)
print("saved", OUT)
